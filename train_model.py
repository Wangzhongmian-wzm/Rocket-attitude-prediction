# -*- coding: utf-8 -*-
"""
训练漂移距离预测模型。

数据来源：楼层抛伞记录.xlsx
  - 11.5 m：Sheet3 中 5 档风速 × 5 次 = 25 个原始实测点（用于拟合 d-v 直线）
  - 1.55 m：三高度汇总中给出的回归结果 k=0.4215, b=0.2633, R²=0.881
            （原始 25 点未在表内，只有回归摘要）
  - 其余高度（7.5 / 9.5 / 13.5 / 15.5 m）均为推算值，不参与训练。

模型：
    d(h, v) = k(h) * v + b(h)
    k(h) = a_k * h^p_k
    b(h) = a_b * h^p_b

输出：
    drift_model.json          模型参数、锚点、不确定性与指标
    data/measured_points.csv  11.5 m 的 25 个实测点
    data/summary_audit.csv    表内各高度数据来源核查
    漂移距离预测表.xlsx        预测结果 / 训练数据 / 模型参数 / 验证指标 / 数据来源
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.model_selection import LeaveOneOut, cross_val_predict
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import PolynomialFeatures

import drift_model

BASE = Path(__file__).resolve().parent
XLSX = BASE / "楼层抛伞记录.xlsx"
OUT_JSON = BASE / "drift_model.json"
OUT_METRICS = BASE / "model_metrics.json"
OUT_XLSX = BASE / "漂移距离预测表.xlsx"
DATA_DIR = BASE / "data"

RNG_SEED = 20261020
N_BOOTSTRAP = 5000
APPLICABLE_HEIGHT = (1.55, 11.5)
APPLICABLE_WIND = (0.2, 3.3)


# --------------------------------------------------------------------------
# 1. 读取数据
# --------------------------------------------------------------------------
def read_sheet3() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """11.5 m 原始实测：5 档风速 × 5 次。返回 v, 25 个 d, 5 个留空时间。"""
    s = pd.read_excel(XLSX, sheet_name="Sheet3", header=None)
    v = s.iloc[2:7, 0].to_numpy(float)
    d = s.iloc[2:7, 1:6].to_numpy(float)
    t = s.iloc[2:7, 8].to_numpy(float)
    if d.shape != (5, 5) or np.isnan(d).any():
        raise ValueError("Sheet3 的 11.5 m 原始数据不完整，请检查表格。")
    return v, d, t


def read_sheet2() -> tuple[np.ndarray, np.ndarray]:
    """Sheet2 中另一版的 25 个原始点（用于稳健性对照，不用于最终训练）。"""
    s = pd.read_excel(XLSX, sheet_name="Sheet2", header=None)
    v, d = [], []
    for block in range(5):
        rows = s.iloc[1 + 5 * block: 6 + 5 * block]
        v.extend(rows.iloc[:, 0].to_numpy(float))
        d.extend(rows.iloc[:, 1].to_numpy(float))
    return np.asarray(v, float), np.asarray(d, float)


def read_summary() -> pd.DataFrame:
    """三高度汇总的 5 行参数：h, k, b, t, R², 来源。"""
    s = pd.read_excel(XLSX, sheet_name="三高度汇总", header=None)
    rows = s.iloc[2:7, 0:6].copy()
    rows.columns = ["h", "k", "b", "t", "r2", "source"]
    rows["r2"] = pd.to_numeric(rows["r2"], errors="coerce")
    return rows.reset_index(drop=True)


def read_summary_table() -> pd.DataFrame:
    """三高度汇总中的推算位移表：h, v=0, v=0.8, v=1.7, v=2.5, v=3.3。"""
    s = pd.read_excel(XLSX, sheet_name="三高度汇总", header=None)
    tab = s.iloc[10:15, 0:6].copy()
    tab.columns = ["h", "v0", "v08", "v17", "v25", "v33"]
    return tab.reset_index(drop=True)


def read_sheet1_audit() -> pd.DataFrame:
    """Sheet1 中各高度 × 风速的均值，加上来源标注。"""
    s = pd.read_excel(XLSX, sheet_name="Sheet1", header=None)
    block = s.iloc[15:40, 1:4].copy()
    block.columns = ["释放高度(m)", "风速(m/s)", "平均漂移距离(m)"]
    block = block.dropna(subset=["释放高度(m)", "风速(m/s)", "平均漂移距离(m)"]).reset_index(drop=True)
    status = []
    for h in block["释放高度(m)"]:
        if abs(h - 11.5) < 1e-9:
            status.append("实测均值（Sheet3 25 个原始点）")
        elif abs(h - 9.5) < 1e-9:
            status.append("来源未标注，未纳入训练")
        else:
            status.append("推算值（非实测，未纳入训练）")
    block["status"] = status
    return block


# --------------------------------------------------------------------------
# 2. 拟合工具
# --------------------------------------------------------------------------
def ols(v: np.ndarray, d: np.ndarray) -> dict:
    """普通最小二乘 d = k v + b，返回系数、协方差与拟合指标。"""
    v = np.asarray(v, float)
    d = np.asarray(d, float)
    X = np.column_stack([v, np.ones_like(v)])
    beta, *_ = np.linalg.lstsq(X, d, rcond=None)
    resid = d - X @ beta
    n = len(d)
    dof = n - 2
    s2 = float(resid @ resid / dof)
    cov = s2 * np.linalg.inv(X.T @ X)
    se = np.sqrt(np.diag(cov))
    ss_res = float(resid @ resid)
    ss_tot = float((d - d.mean()) @ (d - d.mean()))
    r2 = 1.0 - ss_res / ss_tot
    return {
        "n": n,
        "k": float(beta[0]),
        "b": float(beta[1]),
        "se_k": float(se[0]),
        "se_b": float(se[1]),
        "cov": cov,
        "r2": float(r2),
        "r2_adj": float(1.0 - (1.0 - r2) * (n - 1) / dof),
        "rmse": float(math.sqrt(ss_res / n)),
        "mae": float(np.mean(np.abs(resid))),
        "sigma": float(math.sqrt(s2)),
        "resid": resid,
    }


def huber(v: np.ndarray, d: np.ndarray, c: float = 1.345, iters: int = 200) -> tuple[float, float]:
    """Huber 稳健回归，用于敏感性对照。"""
    v = np.asarray(v, float)
    d = np.asarray(d, float)
    X = np.column_stack([v, np.ones_like(v)])
    beta, *_ = np.linalg.lstsq(X, d, rcond=None)
    for _ in range(iters):
        r = d - X @ beta
        scale = 1.4826 * np.median(np.abs(r - np.median(r)))
        if scale <= 0:
            break
        w = np.where(np.abs(r) <= c * scale, 1.0, c * scale / np.abs(r))
        sw = np.sqrt(w)
        new, *_ = np.linalg.lstsq(X * sw[:, None], d * sw, rcond=None)
        if np.max(np.abs(new - beta)) < 1e-10:
            beta = new
            break
        beta = new
    return float(beta[0]), float(beta[1])


def fit_height_law(h1: float, k1: float, b1: float, h2: float, k2: float, b2: float) -> dict:
    """用两个高度锚点确定 k(h)=a_k h^p_k 与 b(h)=a_b h^p_b。"""
    ratio = h2 / h1
    p_k = math.log(k2 / k1) / math.log(ratio)
    a_k = k1 / h1 ** p_k
    p_b = math.log(b2 / b1) / math.log(ratio)
    a_b = b1 / h1 ** p_b
    return {"a_k": a_k, "p_k": p_k, "a_b": a_b, "p_b": p_b}


def anchor_1_55_uncertainty(k: float, b: float, r2: float, n: int,
                            v_design: np.ndarray) -> tuple[float, np.ndarray]:
    """
    根据表内给出的 R²、n 与风速设计，反推 1.55 m 锚点的残差 sigma 和 (k,b) 协方差。
    假设 5 档风速各重复 5 次（表内说明为 25 个原始点）。
    """
    vv = np.tile(np.asarray(v_design, float), 5)
    X = np.column_stack([vv, np.ones_like(vv)])
    sxx = float(((vv - vv.mean()) ** 2).sum())
    ssr = k ** 2 * sxx
    sst = ssr / r2
    sse = sst - ssr
    sigma = math.sqrt(sse / (n - 2))
    cov = sigma ** 2 * np.linalg.inv(X.T @ X)
    return sigma, cov


def bootstrap_anchor_draws(v: np.ndarray, d: np.ndarray, k1: float, b1: float,
                           cov1: np.ndarray, n_draws: int, seed: int) -> tuple[np.ndarray, int]:
    """
    对 11.5 m 的 25 个点做非参数 bootstrap，对 1.55 m 锚点按其抽样协方差做参数抽样。
    拒绝 k<=0 或 b<=0 的不物理样本（漂移距离在 v=0 时不应为负）。
    """
    rng = np.random.default_rng(seed)
    n = len(d)
    draws = []
    rejected = 0
    attempts = 0
    max_attempts = n_draws * 100
    mean1 = np.array([k1, b1], float)
    while len(draws) < n_draws and attempts < max_attempts:
        attempts += 1
        idx = rng.integers(0, n, size=n)
        vb, db = v[idx], d[idx]
        Xb = np.column_stack([vb, np.ones_like(vb)])
        beta, *_ = np.linalg.lstsq(Xb, db, rcond=None)
        k2, b2 = float(beta[0]), float(beta[1])
        k1s, b1s = rng.multivariate_normal(mean1, cov1)
        if k2 > 0 and b2 > 0 and k1s > 0 and b1s > 0:
            draws.append([k1s, b1s, k2, b2])
        else:
            rejected += 1
    if len(draws) < n_draws:
        raise RuntimeError("bootstrap 拒绝率过高，无法生成足够的物理样本。")
    return np.asarray(draws, float), rejected


# --------------------------------------------------------------------------
# 3. 评估指标
# --------------------------------------------------------------------------
def regression_metrics(y: np.ndarray, pred: np.ndarray) -> dict:
    y = np.asarray(y, float)
    pred = np.asarray(pred, float)
    err = y - pred
    ss_res = float(err @ err)
    ss_tot = float((y - y.mean()) @ (y - y.mean()))
    return {
        "n": int(len(y)),
        "rmse": float(math.sqrt(ss_res / len(y))),
        "mae": float(np.mean(np.abs(err))),
        "r2": float(1.0 - ss_res / ss_tot) if ss_tot > 0 else float("nan"),
    }


def loocv_linear(v: np.ndarray, d: np.ndarray) -> tuple[dict, dict]:
    """留一法交叉验证，同时统计 95% 预测区间覆盖率。"""
    n = len(d)
    preds = np.empty(n)
    covered = 0
    widths = []
    for i in range(n):
        mask = np.ones(n, dtype=bool)
        mask[i] = False
        fit = ols(v[mask], d[mask])
        X = np.column_stack([v[mask], np.ones(mask.sum())])
        xtx_inv = np.linalg.inv(X.T @ X)
        x0 = np.array([v[i], 1.0])
        yhat = float(x0 @ np.array([fit["k"], fit["b"]]))
        se = fit["sigma"] * math.sqrt(1.0 + float(x0 @ xtx_inv @ x0))
        tcrit = stats.t.ppf(0.975, n - 3)
        lo, hi = yhat - tcrit * se, yhat + tcrit * se
        preds[i] = yhat
        covered += int(lo <= d[i] <= hi)
        widths.append(hi - lo)
    metrics = regression_metrics(d, preds)
    metrics["pi95_coverage"] = covered / n
    metrics["mean_pi95_width_m"] = float(np.mean(widths))
    return metrics, {"pred": preds.tolist()}


def leave_one_wind_out(v: np.ndarray, d: np.ndarray) -> tuple[list[dict], dict]:
    """留一风速档：用其余风速档拟合，预测被留出档的 5 个点。"""
    out = []
    all_true, all_pred = [], []
    for vv in np.unique(v):
        train = v != vv
        test = v == vv
        fit = ols(v[train], d[train])
        pred = fit["k"] * v[test] + fit["b"]
        m = regression_metrics(d[test], pred)
        out.append({"wind_mps": float(vv), **m,
                    "train_k": fit["k"], "train_b": fit["b"]})
        all_true.extend(d[test].tolist())
        all_pred.extend(pred.tolist())
    pooled = regression_metrics(np.array(all_true), np.array(all_pred))
    return out, pooled


def model_comparison(v: np.ndarray, d: np.ndarray) -> dict:
    """同一高度、仅用风速作输入的若干模型对比（LOOCV）。"""
    X = v.reshape(-1, 1)
    loo = LeaveOneOut()
    candidates = {
        "线性 d = k·v + b": LinearRegression(),
        "二次多项式": make_pipeline(PolynomialFeatures(2, include_bias=False), LinearRegression()),
        "随机森林": RandomForestRegressor(n_estimators=300, min_samples_leaf=2, random_state=42),
        "梯度提升树": GradientBoostingRegressor(n_estimators=300, max_depth=2, learning_rate=0.05,
                                              random_state=42),
    }
    out = {}
    for name, est in candidates.items():
        pred = cross_val_predict(est, X, d, cv=loo)
        out[name] = regression_metrics(d, pred)
    # 同风速组均值的留一基线（利用“风速档”这一离散信息，不能预测新风速）
    group_loo = np.empty_like(d)
    for i in range(len(d)):
        same = (v == v[i])
        same[i] = False
        group_loo[i] = d[same].mean()
    out["同风速组均值（留一基线）"] = regression_metrics(d, group_loo)
    out["全局均值基线"] = regression_metrics(d, np.full_like(d, d.mean()))
    return out


def compare_with_summary_table(tab: pd.DataFrame, params: dict) -> dict:
    """与表内推算位移表逐项对比（一致性检查，不是独立验证）。"""
    winds = [0.0, 0.8, 1.7, 2.5, 3.3]
    rows = []
    for _, row in tab.iterrows():
        h = float(row["h"])
        if h not in (7.5, 11.5, 13.5, 15.5):
            continue
        k = params["a_k"] * h ** params["p_k"]
        b = params["a_b"] * h ** params["p_b"]
        for v, col in zip(winds, ["v0", "v08", "v17", "v25", "v33"]):
            dtab = float(row[col])
            dnew = k * v + b
            rows.append({"h": h, "v": v, "table_d": dtab, "model_d": dnew, "diff": dnew - dtab})
    diffs = [abs(r["diff"]) for r in rows]
    return {
        "n_compared": len(rows),
        "mean_abs_diff_m": float(np.mean(diffs)),
        "max_abs_diff_m": float(np.max(diffs)),
        "details": rows,
    }


def sensitivity_check(v: np.ndarray, d: np.ndarray, tabs: pd.DataFrame, primary: dict,
                      h1: float, k1: float, b1: float) -> dict:
    """对数据版本、稳健回归、高度插值形式做敏感性分析。"""
    grid_h = [7.5, 13.5, 15.5]
    grid_v = [0.8, 1.7, 2.5, 3.3]

    def predict_with(p):
        return np.array([[p["a_k"] * h ** p["p_k"] * vv + p["a_b"] * h ** p["p_b"]
                          for vv in grid_v] for h in grid_h])

    base = predict_with(primary)

    # 1) Sheet2 版本
    v2, d2 = read_sheet2()
    fit2 = ols(v2, d2)
    p2 = fit_height_law(h1, k1, b1, 11.5, fit2["k"], fit2["b"])
    # 2) Huber 稳健回归
    hk, hb = huber(v, d)
    ph = fit_height_law(h1, k1, b1, 11.5, hk, hb)
    # 3) 高度线性插值（不假设幂律）
    def linear_h(h):
        k = k1 + (primary["k_at_h2"] - k1) * (h - h1) / (11.5 - h1)
        b = b1 + (primary["b_at_h2"] - b1) * (h - h1) / (11.5 - h1)
        return k, b

    lin = np.array([[linear_h(h)[0] * vv + linear_h(h)[1] for vv in grid_v] for h in grid_h])

    return {
        "max_abs_diff_sheet2_m": float(np.max(np.abs(predict_with(p2) - base))),
        "max_abs_diff_huber_m": float(np.max(np.abs(predict_with(ph) - base))),
        "max_abs_diff_linear_height_interp_m": float(np.max(np.abs(lin - base))),
        "sheet2_k": fit2["k"],
        "sheet2_b": fit2["b"],
        "huber_k": hk,
        "huber_b": hb,
        "grid_h": grid_h,
        "grid_v": grid_v,
    }


# --------------------------------------------------------------------------
# 4. 输出
# --------------------------------------------------------------------------
def write_excel(pred_df: pd.DataFrame, train_df: pd.DataFrame, params_df: pd.DataFrame,
                metrics_df: pd.DataFrame, audit_df: pd.DataFrame, path: Path) -> None:
    """写出带基本格式的 Excel 工作簿。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pred_df.to_excel(writer, sheet_name="预测结果", index=False)
        train_df.to_excel(writer, sheet_name="训练数据", index=False)
        params_df.to_excel(writer, sheet_name="模型参数", index=False)
        metrics_df.to_excel(writer, sheet_name="验证指标", index=False)
        audit_df.to_excel(writer, sheet_name="数据来源核查", index=False)

    import openpyxl

    wb = openpyxl.load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1C5CAB")
    header_font = Font(color="FFFFFF", bold=True)
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ws.columns:
            max_len = 0
            letter = get_column_letter(col[0].column)
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 60))
            ws.column_dimensions[letter].width = max(10, min(max_len * 1.3 + 2, 55))
        ws.row_dimensions[1].height = 30

    ws = wb["预测结果"]
    headers = [c.value for c in ws[1]]
    if "适用性" in headers:
        idx = headers.index("适用性")
        interp_fill = PatternFill("solid", fgColor="EAF2FB")
        double_fill = PatternFill("solid", fgColor="F4CCCC")
        for row in ws.iter_rows(min_row=2):
            status = str(row[idx].value or "")
            if "双重外推" in status:
                fill = double_fill
            elif "外推" in status:
                fill = warn_fill
            elif "内插" in status:
                fill = interp_fill
            else:
                fill = None
            if fill is not None:
                for cell in row:
                    cell.fill = fill

    # 数字格式
    ws = wb["预测结果"]
    headers = [c.value for c in ws[1]]
    fmt = {
        "释放高度(m)": "0.00", "风速(m/s)": "0.00", "预测漂移距离(m)": "0.00",
        "95%下限(m)": "0.00", "95%上限(m)": "0.00", "留空时间(s)": "0.00",
        "k(s)": "0.000", "b(m)": "0.000",
    }
    for name, code in fmt.items():
        if name in headers:
            j = headers.index(name) + 1
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                row[0].number_format = code

    # 验证指标：统一小数位，覆盖率显示为百分比
    ws = wb["验证指标"]
    for row in ws.iter_rows(min_row=2):
        label, value = row[0].value, row[1].value
        if label == "95% 区间留一覆盖率":
            row[1].number_format = "0.0%"
        elif isinstance(value, (int, float)):
            row[1].number_format = "0.0000"

    wb.properties.title = "漂移距离预测表"
    wb.properties.creator = "train_model.py"
    wb.save(path)


# --------------------------------------------------------------------------
# 5. 主流程
# --------------------------------------------------------------------------
def main() -> None:
    DATA_DIR.mkdir(exist_ok=True)

    # --- 数据 ---
    v, dmat, t_obs = read_sheet3()
    V = np.repeat(v, 5)
    D = dmat.reshape(-1)
    summary = read_summary()
    summary_tab = read_summary_table()

    anchor_row = summary[summary["h"] == 1.55].iloc[0]
    H1, K1, B1, T1, R2_1 = 1.55, float(anchor_row["k"]), float(anchor_row["b"]), float(anchor_row["t"]), float(anchor_row["r2"])
    H2, T2 = 11.5, float(summary[summary["h"] == 11.5].iloc[0]["t"])

    fit = ols(V, D)
    K2, B2 = fit["k"], fit["b"]
    params = fit_height_law(H1, K1, B1, H2, K2, B2)
    params["k_at_h2"] = K2
    params["b_at_h2"] = B2

    # 留空时间 t(h) 的幂律（用 1.55 / 11.5 两个实测锚点）
    p_t = math.log(T2 / T1) / math.log(H2 / H1)
    a_t = T1 / H1 ** p_t

    # --- 不确定性 ---
    sigma1, cov1 = anchor_1_55_uncertainty(K1, B1, R2_1, 25, v)
    draws, rejected = bootstrap_anchor_draws(V, D, K1, B1, cov1, N_BOOTSTRAP, RNG_SEED)

    # --- 评估 ---
    loocv, _ = loocv_linear(V, D)
    lowo, lowo_pooled = leave_one_wind_out(V, D)
    comparison = model_comparison(V, D)
    consistency = compare_with_summary_table(summary_tab, params)
    sens = sensitivity_check(V, D, summary_tab, params, H1, K1, B1)

    # 同风速组内的测量噪声（用于解释 LOOCV 误差下限）
    group_means = np.array([D[V == vv].mean() for vv in V])
    noise_floor = regression_metrics(D, group_means)

    metrics = {
        "fit_11_5m": {k: fit[k] for k in
                      ["n", "k", "b", "se_k", "se_b", "r2", "r2_adj", "rmse", "mae", "sigma"]},
        "anchor_1_55m": {
            "n": 25, "k": K1, "b": B1, "r2": R2_1,
            "sigma_est": sigma1,
            "se_k_est": float(math.sqrt(cov1[0, 0])),
            "se_b_est": float(math.sqrt(cov1[1, 1])),
            "note": "表内只有回归摘要，原始 25 点未提供；SE 由 R² 与 5×5 风速设计反推。",
        },
        "loocv_11_5m": loocv,
        "leave_one_wind_speed_out": {"details": lowo, "pooled": lowo_pooled},
        "model_comparison_loocv": comparison,
        "within_group_noise_floor": noise_floor,
        "table_consistency": consistency,
        "sensitivity": sens,
        "bootstrap": {
            "n_draws": int(len(draws)),
            "rejected_nonpositive": int(rejected),
            "seed": RNG_SEED,
        },
    }

    # --- 模型 JSON ---
    model = {
        "schema_version": 1,
        "trained_at": "2026-10-20",
        "target": "漂移距离 d (m)",
        "inputs": ["释放高度 h (m)", "风速 v (m/s)"],
        "formula": "d(h,v) = a_k * h**p_k * v + a_b * h**p_b",
        "parameters": {
            "a_k": params["a_k"], "p_k": params["p_k"],
            "a_b": params["a_b"], "p_b": params["p_b"],
            "a_t": a_t, "p_t": p_t,
        },
        "anchors": {
            "h1": H1, "k1": K1, "b1": B1, "t1": T1,
            "h2": H2, "k2": K2, "b2": B2, "t2": T2,
            "height_ratio": H2 / H1,
            "k2_source": "Sheet3 的 25 个原始点 OLS 重拟合",
            "k1_source": "三高度汇总给出的 1.55 m 回归系数（原始点未提供）",
        },
        "uncertainty": {
            "h1": H1, "h2": H2,
            "sigma_h1": sigma1, "sigma_h2": fit["sigma"],
            "sigma_note": "sigma_h1 由表内 R²=0.881 与 5×5 风速设计估算；sigma_h2 为 25 点 OLS 残差标准差。",
            "n_draws": int(len(draws)),
            "rejected_nonpositive": int(rejected),
            "draws_note": "每行 [k(1.55), b(1.55), k(11.5), b(11.5)]；已剔除非物理的非正样本。",
            "anchor_draws": np.round(draws, 8).tolist(),
        },
        "applicable_range": {"height": list(APPLICABLE_HEIGHT), "wind": list(APPLICABLE_WIND)},
        "metrics": metrics,
        "data_provenance": {
            "measured_heights": [1.55, 11.5],
            "measured_points": {"1.55": 25, "11.5": 25},
            "projected_heights": [7.5, 13.5, 15.5],
            "undocumented_height": [9.5],
            "excluded_from_training": "7.5 / 9.5 / 13.5 / 15.5 m 均为推算或来源未标注，不参与训练。",
        },
        "limitations": [
            "只有两个高度的实测锚点，且 1.55 m 是课内风扇装置、11.5 m 是室外抛伞装置，两者并非同一装置。",
            "1.55 m 只有回归摘要，没有 25 个原始点；其不确定度由 R² 反推。",
            "13.5 / 15.5 m 超出 11.5 m 实测锚点，属于高度外推。",
            "95% 区间只反映锚点抽样不确定性与测量散差，不包含模型形式错误；外推时真实误差可能更大。",
            "95% 区间是单次抛掷的预测区间；下限已按“漂移距离非负”截断到 0。",
            "风速与漂移方向、伞衣充气状态、载荷质量、温湿度、地形等因素未纳入模型。",
        ],
    }

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(model, f, ensure_ascii=False, indent=2)
    with open(OUT_METRICS, "w", encoding="utf-8") as f:
        json.dump(metrics, f, ensure_ascii=False, indent=2)

    # --- CSV ---
    train_df = pd.DataFrame({
        "序号": np.arange(1, 26),
        "释放高度(m)": 11.5,
        "风速(m/s)": V,
        "漂移距离(m)": D,
        "留空时间(s)": np.repeat(t_obs, 5),
        "来源": "Sheet3 11.5 m 实测",
    })
    train_df.to_csv(DATA_DIR / "measured_points.csv", index=False, encoding="utf-8-sig")

    audit_df = read_sheet1_audit()
    audit_df.to_csv(DATA_DIR / "summary_audit.csv", index=False, encoding="utf-8-sig")

    # --- 预测网格 ---
    pred_heights = [1.55, 7.5, 9.5, 11.5, 13.5, 15.5]
    pred_winds = [0.0, 0.2, 0.8, 1.7, 2.5, 3.3]
    rows = []
    for h in pred_heights:
        for vv in pred_winds:
            r = drift_model.predict(h, vv, model=model)
            rows.append({
                "释放高度(m)": h,
                "风速(m/s)": vv,
                "预测漂移距离(m)": r["drift_m"],
                "95%下限(m)": r["lo95"],
                "95%上限(m)": r["hi95"],
                "留空时间(s)": r["hang_time_s"],
                "k(s)": r["k_s"],
                "b(m)": r["b_m"],
                "适用性": r["applicability"],
                "提示": "；".join(r["warnings"]) if r["warnings"] else "",
            })
    pred_df = pd.DataFrame(rows)
    pred_df.to_csv(BASE / "预测结果.csv", index=False, encoding="utf-8-sig")

    # 模型参数表
    param_rows = [
        ("模型形式", "d(h,v) = k(h)·v + b(h)"),
        ("k(h)", f"{params['a_k']:.6f} · h^{params['p_k']:.6f}"),
        ("b(h)", f"{params['a_b']:.6f} · h^{params['p_b']:.6f}"),
        ("a_k", params["a_k"]), ("p_k", params["p_k"]),
        ("a_b", params["a_b"]), ("p_b", params["p_b"]),
        ("k(1.55 m)", K1), ("b(1.55 m)", B1),
        ("k(11.5 m)", K2), ("b(11.5 m)", B2),
        ("t(h)", f"{a_t:.6f} · h^{p_t:.6f}"),
        ("训练样本", "25 个实测点（11.5 m）+ 1 个回归锚点（1.55 m）"),
        ("适用高度", f"{APPLICABLE_HEIGHT[0]} – {APPLICABLE_HEIGHT[1]} m（超出为外推）"),
        ("适用风速", f"{APPLICABLE_WIND[0]} – {APPLICABLE_WIND[1]} m/s（超出为外推）"),
        ("95% 区间", "锚点抽样 + 残差噪声的 2.5%–97.5% 分位数"),
        ("模型文件", "drift_model.json"),
    ]
    param_df = pd.DataFrame(param_rows, columns=["项目", "数值"])

    # 验证指标表
    metric_rows = [
        ("11.5 m 拟合 R²", fit["r2"]),
        ("11.5 m 拟合 RMSE (m)", fit["rmse"]),
        ("11.5 m 留一法 R²", loocv["r2"]),
        ("11.5 m 留一法 RMSE (m)", loocv["rmse"]),
        ("11.5 m 留一法 MAE (m)", loocv["mae"]),
        ("95% 区间留一覆盖率", loocv["pi95_coverage"]),
        ("留一风速档合并 RMSE (m)", lowo_pooled["rmse"]),
        ("同风速组内散差 RMSE (m)", noise_floor["rmse"]),
        ("与表内推算法的平均差 (m)", consistency["mean_abs_diff_m"]),
        ("与表内推算法的最大差 (m)", consistency["max_abs_diff_m"]),
        ("Sheet2 版本敏感性 (m)", sens["max_abs_diff_sheet2_m"]),
        ("Huber 稳健回归敏感性 (m)", sens["max_abs_diff_huber_m"]),
        ("高度线性插值敏感性 (m)", sens["max_abs_diff_linear_height_interp_m"]),
    ]
    metrics_df = pd.DataFrame(metric_rows, columns=["指标", "数值"])
    metrics_df["数值"] = metrics_df["数值"].map(lambda x: round(x, 6) if isinstance(x, float) else x)

    write_excel(pred_df, train_df, param_df, metrics_df, audit_df, OUT_XLSX)

    # --- 终端摘要 ---
    print("=" * 72)
    print("训练完成")
    print("=" * 72)
    print(f"11.5 m 原始点: n=25, k={K2:.4f} (SE {fit['se_k']:.4f}), b={B2:.4f} (SE {fit['se_b']:.4f})")
    print(f"               R²={fit['r2']:.4f}, RMSE={fit['rmse']:.4f} m, sigma={fit['sigma']:.4f} m")
    print(f"1.55 m 锚点:   k={K1:.4f}, b={B1:.4f}, R²={R2_1:.3f} (原始点未提供)")
    print(f"高度幂律:      k(h) = {params['a_k']:.4f} · h^{params['p_k']:.4f}")
    print(f"               b(h) = {params['a_b']:.4f} · h^{params['p_b']:.4f}")
    print(f"留空时间:      t(h) = {a_t:.4f} · h^{p_t:.4f}")
    print(f"LOOCV:         R²={loocv['r2']:.4f}, RMSE={loocv['rmse']:.4f} m, "
          f"MAE={loocv['mae']:.4f} m, PI95 覆盖率={loocv['pi95_coverage']:.1%}")
    print(f"留一风速档:    合并 RMSE={lowo_pooled['rmse']:.4f} m; "
          f"同风速组噪声下限={noise_floor['rmse']:.4f} m")
    print(f"与表内推算值:  平均差={consistency['mean_abs_diff_m']:.3f} m, "
          f"最大差={consistency['max_abs_diff_m']:.3f} m")
    print(f"bootstrap:     {len(draws)} 组抽样, 拒绝非正样本 {rejected} 组")
    print()
    print("示例：h=12 m, v=2.0 m/s")
    print(drift_model.format_result(drift_model.predict(12.0, 2.0, model=model)))
    print()
    print(f"已保存: {OUT_JSON.name}, {OUT_METRICS.name}, {OUT_XLSX.name}, "
          f"预测结果.csv, data/*.csv")


if __name__ == "__main__":
    main()
